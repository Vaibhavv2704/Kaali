"""Publish original 2024 totals/categories and documented approximate references."""
import hashlib
import json
from pathlib import Path
import openpyxl
from shapely.geometry import shape
from shapely.ops import unary_union
from count_model import count
from train_count_model import observations, independent_report_check

ROOT = Path(__file__).resolve().parent
POLICE = 'https://yuva.delhipolice.gov.in/contact-us.html'
ANCHORS = {
    'Central': ('node/12441175891', 'Office Of The Deputy Commissioner Of Police Central District'),
    'East': ('node/11959842021', 'Police Station Madhu Vihar'),
    'New Delhi': ('node/5634522760', 'Police Station Mandir Marg'),
    'North': ('node/13514326623', 'Civil Lines Police Station'),
    'North-East': ('node/663758609', 'Seelampur'),
    'North-West': ('node/1402109581', 'Ashok Vihar Police Station'),
    'Outer': ('node/6465991447', 'Mundka'),
    'Outer North': ('node/762031782', 'Bawana Police station'),
    'Rohini': ('node/663766390', 'Rohini East'),
    'Dwarka': ('node/2408446621', 'Dwarka Police Station'),
    'Shahdara': ('node/12596382943', 'DCP Shahdara District Office'),
    'South': ('node/2004001910', 'Police Station Malviya Nagar'),
    'South-East': ('node/7099818528', 'Police Station Badarpur'),
    'South-West': ('node/2832938196', 'Vasant Vihar'),
    'West': ('node/2445916782', 'Punjabi Bagh Police Station'),
}
HIGHLIGHTS = {3: 'Rape', 11: 'Assault with intent to outrage modesty',
              41: 'Kidnapping / abduction (source group)', 29: 'Dowry deaths',
              30: 'Cruelty by husband / relatives', 59: 'POCSO · girl-child cases'}


def main():
    config = json.loads((ROOT/'config/counts-v1.json').read_text())
    observations_, checks, _ = observations(config)
    independent_report_check(checks)
    if any(c['difference'] != 0 or c['independent_difference'] != 0 for c in checks):
        raise ValueError('State totals do not reconcile')
    source = next(s for s in config['sources'] if s['year'] == 2024)
    path = ROOT/source['path']
    if hashlib.sha256(path.read_bytes()).hexdigest() != source['sha256']:
        raise ValueError('Source changed')
    sheet = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    help_ = json.loads((ROOT.parent/'public/data/delhi-ncr/help.geojson').read_text())
    osm_receipt = next(s for s in json.loads((ROOT/'config/phase1-inputs.json').read_text())['inputs'] if s['id'] == 'osm-help')
    raw_osm = ROOT/osm_receipt['path']
    if hashlib.sha256(raw_osm.read_bytes()).hexdigest() != osm_receipt['sha256']:
        raise ValueError('OSM source snapshot changed; reference review required')
    raw_features = json.loads(raw_osm.read_text())['features']
    boundary = json.loads((ROOT.parent/'public/data/delhi-ncr/delhi-boundary.geojson').read_text())
    delhi_outline = unary_union([shape(f['geometry']) for f in boundary['features']])
    records = []
    for row in observations_:
        if row['city'] != 'delhi' or row['year'] != 2024:
            continue
        unit = row['unit']; key, name = ANCHORS[unit]
        matches = [f for f in help_['features'] if f['properties']['id'] == key and f['properties']['name'] == name]
        if len(matches) != 1 or matches[0]['geometry']['type'] != 'Point':
            raise ValueError('Reviewed reference identity changed: '+unit)
        point = matches[0]
        raw_match = [f for f in raw_features if f['properties']['id'] == key and f['properties']['name'] == name]
        if len(raw_match) != 1 or raw_match[0]['geometry'] != point['geometry']:
            raise ValueError('Published reference differs from retrieved OSM snapshot')
        if not delhi_outline.covers(shape(point['geometry'])):
            raise ValueError('Reference is outside Delhi NCT; do not assign cross-border points')
        cells = list(sheet.values)[row['source_row']-1]
        categories = []; parent = ''
        for column in range(3, 67):
            primary = sheet.cell(2, column).value
            secondary = sheet.cell(3, column).value
            if primary: parent = str(primary).strip()
            label = parent + (' · '+str(secondary).strip() if secondary else '')
            categories.append({'column': column, 'label': label, 'count': count(cells[column-1]),
                               'kind': 'subtotal' if 'Total' in label or 'total' in label else 'source_heading',
                               'highlight': HIGHLIGHTS.get(column)})
        records.append({'id': unit, 'name': unit, 'year': 2024, 'count': row['count'],
            'countType': 'recorded_total', 'sourceRow': row['source_row'], 'categories': categories,
            'missingColumns': [c['column'] for c in categories if c['count'] is None],
            'categoryMismatches': row['category_mismatches'],
            'reference': {'coordinates': point['geometry']['coordinates'], 'name': name,
                'kind': point['properties']['kind'], 'source': point['properties']['source'],
                'assignmentSource': POLICE, 'status': 'approximate_reference',
                'reviewedAt': '2026-10-01'}})
    if len(records) != 15: raise ValueError('Expected 15 geographic police districts')
    receipt = json.loads((ROOT/'sources/delhi-district-2024-receipt.json').read_text())
    pdf = next(s for s in receipt['sources'] if s['id'] == 'ncrb-2024-volume-1')
    result = {'schemaVersion': 1, 'year': 2024, 'sourceUrl': pdf['officialCatalogUrl'],
        'reportUrl': pdf['downloadUrl'], 'workbookSha256': source['sha256'],
        'provenanceNote': 'Existing original-titled district workbook; original download URL/date unavailable. Explicit totals reconcile with NCRB Table 3A.1. Government publication reuse terms require review.',
        'stateTotal': 13396, 'geographicTotal': sum(r['count'] for r in records),
        'specialUnitTotal': 13396-sum(r['count'] for r in records),
        'referenceNote': 'Approximate public-place reference coordinates from OSM, linked to district stations/localities listed by Delhi Police; not incident locations, centroids, historical boundaries or danger radii.',
        'records': sorted(records, key=lambda r: r['name'])}
    output = ROOT.parent/'public/data/delhi-historical-districts.json'
    output.write_text(json.dumps(result, indent=2, allow_nan=False)+'\n', encoding='utf-8')
    print('Published 15 historical districts; measured totals, no safety scores or invented geometry.')


if __name__ == '__main__': main()
